"""
IntentCloud - 150+ Multi-Format Corpus Generator & Expander (Objective 1)
Generates 150+ mixed-format documents (PDF, DOCX, TXT) across 10 technical domains:
1. Distributed Systems (Kafka, Raft, Paxos, Stream Processing)
2. Software Architecture (Microservices, Event-Driven, Circuit Breakers, CQRS)
3. Machine Learning & Deep Learning (Transformers, Attention, CNNs, Optimization)
4. Information Retrieval & Search (BM25, Dense Embeddings, RRF, Cross-Encoders)
5. Cloud Computing & Kubernetes (Containers, Orchestration, Ingress, Pod Scaling)
6. Edge & IoT Systems (ARM Cortex, Raspberry Pi, Thermal Profiles, MQTT)
7. Database Systems & Storage (LSM Trees, B-Trees, Write-Ahead Logs, Sharding)
8. Cybersecurity & Zero-Trust (TLS, Cryptography, OAuth2, Tunnels)
9. Operating Systems & Concurrency (Virtual Memory, Scheduling, Mutexes, Async)
10. DevOps & SRE (CI/CD, Monitoring, Observability, SLO/SLA)

Supports version lineage: generates v1 and v2 pairs to test Objective 3 recency/lineage.
"""

from pathlib import Path
import os
import fitz  # PyMuPDF
from docx import Document

BASE_DIR = Path(__file__).resolve().parent.parent / "test_corpus" / "corpus_150"

DOMAINS = {
    "distributed_systems": [
        ("kafka_replication_protocols", "Apache Kafka partition replication, in-sync replicas (ISR), and leader election under network partitions."),
        ("stream_processing_backpressure", "Reactive streams backpressure handling, flow control buffers, and dynamic rate limiting."),
        ("raft_consensus_algorithm", "Raft consensus protocol state transitions: leader election, log replication, and commit indexes."),
        ("distributed_transactions_2pc", "Two-Phase Commit (2PC) and Saga orchestration for distributed data consistency."),
        ("vector_clocks_and_causality", "Logical clocks, Lamport timestamps, and vector clocks for partial ordering in decentralized systems."),
    ],
    "software_architecture": [
        ("microservices_circuit_breaker", "Resilience4j and Netflix Hystrix circuit breaker states: closed, open, half-open failure thresholds."),
        ("event_driven_event_sourcing", "Event sourcing architecture with CQRS, event stores, and snapshotting for aggregate reconstruction."),
        ("api_gateway_rate_limiting", "Token bucket and leaky bucket rate limiting algorithms implemented in Envoy and Kong API Gateways."),
        ("domain_driven_design_patterns", "Bounded contexts, aggregates, entities, and value objects in enterprise domain-driven design."),
        ("database_sharding_strategies", "Horizontal database partitioning, consistent hashing, and cross-shard querying trade-offs."),
    ],
    "machine_learning": [
        ("transformer_multihead_attention", "Scaled dot-product attention and multi-head attention projection mechanics in transformer encoders."),
        ("convolutional_neural_networks", "Convolution kernels, receptive fields, stride, padding, and feature map pooling in deep CNNs."),
        ("adam_optimizer_mathematics", "First and second moment estimation with bias corrections in Adam and AdamW optimizers."),
        ("quantization_methods_int8", "Post-training quantization (PTQ) and quantization-aware training (QAT) mapping FP32 weights to INT8."),
        ("loss_functions_and_backprop", "Cross-entropy, negative log-likelihood, and backpropagation chain rule calculus."),
    ],
    "information_retrieval": [
        ("bm25_probabilistic_formulation", "Okapi BM25 scoring formula: inverse document frequency (IDF) and sub-linear term frequency scaling."),
        ("dense_passage_retrieval_biencoder", "Dual-encoder architectures for dense passage retrieval using inner product nearest neighbors."),
        ("reciprocal_rank_fusion_mechanics", "Scale-invariant Reciprocal Rank Fusion (RRF, k=60) combining heterogeneous ranking distributions."),
        ("cross_encoder_reranking_attention", "Full cross-attention token interactions in cross-encoders for precision query-document reranking."),
        ("calibrated_confidence_abstention", "Logistic sigmoid confidence gating on cross-encoder logits for out-of-domain query abstention."),
    ],
    "cloud_computing": [
        ("kubernetes_pod_lifecycle", "Kubernetes pod phases, init containers, readiness probes, liveness probes, and graceful termination."),
        ("container_runtime_cgroups", "Linux cgroups v2, namespaces, and OCI container isolation mechanics under containerd."),
        ("serverless_cold_start_mitigation", "Function-as-a-Service cold start latency analysis, pre-warming, and snapshot restore."),
        ("service_mesh_mtls_security", "Istio and Linkerd service mesh architectures with automatic mTLS certificate rotation."),
        ("cloud_cost_optimization_sre", "Spot instance orchestration, auto-scaling horizontal pod autoscalers (HPA), and resource quotas."),
    ],
    "edge_iot": [
        ("raspberry_pi_thermal_throttling", "Thermal dissipation on ARM Cortex-A72: passive vs active cooling and CPU frequency governor scaling."),
        ("mqtt_protocol_iot_telemetry", "MQTT QoS levels 0, 1, 2, keep-alive pinging, and topic wildcards for constrained edge devices."),
        ("embedded_vector_databases_edge", "Embedded vector storage using Qdrant and SQLite-VSS on memory-constrained edge hardware."),
        ("edge_inference_onnx_runtime", "Optimizing neural inference latency on ARM NEON SIMD instructions using ONNX Runtime."),
        ("low_power_energy_characterization", "Measuring Joules-per-query on 15-watt single-board computers under continuous inference loads."),
    ],
    "database_systems": [
        ("lsm_tree_storage_engines", "Log-Structured Merge-tree architecture: MemTable, Write-Ahead Log (WAL), and SSTable compaction."),
        ("b_tree_indexing_internals", "B+ Tree node splitting, branch factor, disk page alignment, and range scan performance."),
        ("acid_isolation_levels", "Read Uncommitted, Read Committed, Repeatable Read, and Serializable snapshot isolation anomalies."),
        ("write_ahead_logging_recovery", "ARIES recovery algorithm: analysis, redo, and undo phases for database crash consistency."),
        ("columnar_storage_parquet", "Column-oriented storage formats, dictionary encoding, run-length encoding, and bit-packing in Parquet."),
    ],
    "cybersecurity": [
        ("tls_handshake_cryptography", "TLS 1.3 cryptographic handshake: Diffie-Hellman key exchange, AEAD ciphers, and 0-RTT risks."),
        ("oauth2_pkce_authorization", "OAuth 2.0 Authorization Code flow with Proof Key for Code Exchange (PKCE) for native clients."),
        ("zero_trust_network_tunnels", "Outbound-only reverse tunnels, identity-aware proxies, and Cloudflare Zero-Trust architecture."),
        ("cross_site_request_forgery", "SameSite cookie policies, anti-CSRF token verification, and defense-in-depth web protection."),
        ("sql_injection_defense_mechanisms", "Parameterized queries, abstract syntax tree (AST) validation, and least-privilege database roles."),
    ],
    "operating_systems": [
        ("virtual_memory_paging_algorithms", "Page tables, Translation Lookaside Buffer (TLB) misses, and Least Recently Used (LRU) page replacement."),
        ("process_scheduling_algorithms", "Completely Fair Scheduler (CFS), red-black tree runqueues, and priority decay in the Linux kernel."),
        ("concurrency_deadlock_detection", "Coffman conditions, resource allocation graphs, and Banker's algorithm for deadlock avoidance."),
        ("async_event_loop_epoll", "Non-blocking I/O multiplexing with epoll, kqueue, and asynchronous event loops in modern runtimes."),
        ("memory_allocator_ptmalloc_jemalloc", "Thread caching, arenas, heap fragmentation, and slab allocation mechanics in jemalloc."),
    ],
    "devops_sre": [
        ("ci_cd_pipeline_best_practices", "Automated regression testing, semantic versioning, and canary deployment verification in GitHub Actions."),
        ("prometheus_metrics_monitoring", "Counters, gauges, histograms, and summary metric types in Prometheus time-series databases."),
        ("distributed_tracing_opentelemetry", "Span context propagation, W3C traceparent headers, and distributed trace sampling with Jaeger."),
        ("error_budgets_and_slo", "Defining Service Level Indicators (SLI), Service Level Objectives (SLO), and error budget burn rate alerts."),
        ("infrastructure_as_code_terraform", "Declarative state management, resource dependency graphs, and drift detection in Terraform."),
    ],
}

def generate_pdf(file_path: Path, title: str, content: str):
    """Create a formatted PDF document using PyMuPDF."""
    doc = fitz.open()
    page = doc.new_page(width=595, height=842) # A4
    text = (
        f"{title.upper().replace('_', ' ')}\n"
        f"{'='*50}\n\n"
        f"Domain: {file_path.parent.name.replace('_', ' ').title()}\n"
        f"Document ID: {file_path.stem}\n\n"
        f"Executive Summary:\n{content}\n\n"
        f"Technical Specifications:\n"
        f"This document provides rigorous theoretical formulation and operational guidelines "
        f"for engineering and research deployment within the IntentCloud knowledge corpus.\n\n"
        f"Detailed Section:\n"
        f"1. Core Principles: {content}\n"
        f"2. Methodological Architecture: High-throughput ingestion, dual-stream indexing, and "
        f"scale-invariant Reciprocal Rank Fusion.\n"
        f"3. Verification: Complies with IEEE standard evaluation benchmarks.\n"
    )
    rect = fitz.Rect(50, 50, 545, 792)
    page.insert_textbox(rect, text, fontsize=11, fontname="helv", align=0)
    doc.save(str(file_path))
    doc.close()

def generate_docx(file_path: Path, title: str, content: str):
    """Create a formatted DOCX document using python-docx."""
    doc = Document()
    doc.add_heading(title.replace('_', ' ').title(), level=1)
    p_meta = doc.add_paragraph()
    p_meta.add_run(f"Domain: {file_path.parent.name.replace('_', ' ').title()}\n").bold = True
    p_meta.add_run(f"Document File: {file_path.name}\n")
    doc.add_heading("Overview & Synopsis", level=2)
    doc.add_paragraph(content)
    doc.add_heading("Operational Details", level=2)
    doc.add_paragraph(
        "This specification is catalogued in IntentCloud for intent-aware hybrid retrieval, "
        "supporting query-adaptive recency scoring and multi-graded evaluation."
    )
    doc.save(str(file_path))

def generate_txt(file_path: Path, title: str, content: str):
    """Create a plain text document."""
    text = (
        f"# {title.replace('_', ' ').title()}\n\n"
        f"Domain: {file_path.parent.name.replace('_', ' ').title()}\n"
        f"Filename: {file_path.name}\n\n"
        f"## Abstract\n{content}\n\n"
        f"## Implementation Notes\n"
        f"Local-first document indexed into 384-dimensional dense semantic vectors "
        f"and 1,000,003-dimensional sparse token hashes.\n"
    )
    with open(file_path, "w", encoding="utf-8") as f:
        f.write(text)

def build_corpus_150():
    """Build the complete 150+ document corpus across PDF, DOCX, and TXT formats."""
    BASE_DIR.mkdir(parents=True, exist_ok=True)
    count = 0

    for domain_name, topic_list in DOMAINS.items():
        domain_dir = BASE_DIR / domain_name
        domain_dir.mkdir(parents=True, exist_ok=True)
        
        for topic_slug, topic_desc in topic_list:
            # 1. Base PDF version (v1)
            pdf_path_v1 = domain_dir / f"{topic_slug}_v1.pdf"
            generate_pdf(pdf_path_v1, f"{topic_slug} (Initial Draft)", topic_desc)
            count += 1

            # 2. Updated DOCX version (v2 - Lineage pair)
            docx_path_v2 = domain_dir / f"{topic_slug}_v2.docx"
            generate_docx(docx_path_v2, f"{topic_slug} (Revised Specification)", f"Revised edition: {topic_desc} Includes updated benchmarks and latency parameters.")
            count += 1

            # 3. Plain text operational notes (v1 or standalone)
            txt_path = domain_dir / f"{topic_slug}_notes.txt"
            generate_txt(txt_path, f"{topic_slug} Notes", f"Operational engineering notes: {topic_desc}")
            count += 1

    print(f"✓ Successfully generated {count} mixed-format documents in {BASE_DIR}")
    return count

if __name__ == "__main__":
    build_corpus_150()
